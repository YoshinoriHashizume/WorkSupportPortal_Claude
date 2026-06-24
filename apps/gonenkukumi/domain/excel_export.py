from __future__ import annotations

import colorsys
from io import BytesIO
from numbers import Number
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


THIN_GRAY = Side(style="thin", color="94A3B8")
BLACK = Side(style="medium", color="334155")
GRID_BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)
SECTION_BORDER = Border(left=BLACK, right=BLACK, top=BLACK, bottom=BLACK)
HSL_PATTERN = re.compile(r"hsl\((?P<hue>[\d.]+)\s+(?P<saturation>[\d.]+)%\s+(?P<lightness>[\d.]+)%\)")
PAD_DAY_COLOR = "B8C4D0"


def display_number(value: object) -> object:
    if isinstance(value, Number) and value == 0:
        return None
    return value


def display_year_month(value: object) -> str:
    return str(value or "").replace("-", "/")


def display_date(value: object) -> str:
    return str(value or "").replace("-", "/")


def display_width(value: object) -> int:
    if value is None:
        return 0
    text = str(value)
    width = 0
    for char in text:
        if char == "\n":
            continue
        width += 2 if unicodedata.east_asian_width(char) in {"F", "W", "A"} else 1
    return width


def fit_column_widths(sheet) -> None:
    merged_cells = {
        (row, col)
        for merged_range in sheet.merged_cells.ranges
        for row in range(merged_range.min_row, merged_range.max_row + 1)
        for col in range(merged_range.min_col, merged_range.max_col + 1)
    }
    for col in range(1, sheet.max_column + 1):
        max_width = 0
        for row in range(1, sheet.max_row + 1):
            if (row, col) in merged_cells:
                continue
            max_width = max(max_width, display_width(sheet.cell(row=row, column=col).value))
        column = sheet.column_dimensions[get_column_letter(col)]
        column.bestFit = True
        column.width = max(max_width + 2, 4)


def excel_color(value: object, fallback: str) -> str:
    text = str(value or "").strip()
    if text.startswith("#") and len(text) == 7:
        return text[1:].upper()

    hsl_match = HSL_PATTERN.fullmatch(text)
    if hsl_match:
        hue = float(hsl_match.group("hue")) / 360
        saturation = float(hsl_match.group("saturation")) / 100
        lightness = float(hsl_match.group("lightness")) / 100
        red, green, blue = colorsys.hls_to_rgb(hue, lightness, saturation)
        return f"{round(red * 255):02X}{round(green * 255):02X}{round(blue * 255):02X}"

    return fallback


def apply_range_fill(sheet, row: int, start_col: int, end_col: int, color: str) -> None:
    fill = PatternFill("solid", fgColor=color)
    for col in range(start_col, end_col + 1):
        sheet.cell(row=row, column=col).fill = fill


def apply_outer_border(sheet, start_row: int, start_col: int, end_row: int, end_col: int) -> None:
    is_single_row_merged_outline = start_row == end_row
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            cell = sheet.cell(row=row, column=col)
            current = cell.border
            cell.border = Border(
                left=BLACK if col == start_col else current.left,
                right=BLACK if col == end_col or (is_single_row_merged_outline and col == start_col) else current.right,
                top=BLACK if row == start_row else current.top,
                bottom=BLACK if row == end_row or (is_single_row_merged_outline and col == start_col) else current.bottom,
            )


def apply_horizontal_border(sheet, row: int, start_col: int, end_col: int) -> None:
    side = Side(style="thin", color="334155")
    for col in range(start_col, end_col + 1):
        top_cell = sheet.cell(row=row, column=col)
        top_border = top_cell.border
        top_cell.border = Border(
            left=top_border.left,
            right=top_border.right,
            top=top_border.top,
            bottom=side,
        )

        bottom_cell = sheet.cell(row=row + 1, column=col)
        bottom_border = bottom_cell.border
        bottom_cell.border = Border(
            left=bottom_border.left,
            right=bottom_border.right,
            top=side,
            bottom=bottom_border.bottom,
        )


def block_theme_colors(block: dict[str, object]) -> dict[str, str]:
    theme = block.get("theme") if isinstance(block.get("theme"), dict) else {}
    return {
        "bg": excel_color(theme.get("bg"), "DCFCE7"),
        "lightBg": excel_color(theme.get("lightBg"), "F0FDF4"),
        "veryLightBg": excel_color(theme.get("veryLightBg"), "F7FEF9"),
        "border": excel_color(theme.get("border"), "86EFAC"),
        "outerBorder": excel_color(theme.get("outerBorder"), "22C55E"),
    }


def panel_active_days(panel: dict[str, object]) -> int:
    try:
        return int(panel.get("activeDays") or 31)
    except (TypeError, ValueError):
        return 31


def panel_holiday_days(panel: dict[str, object]) -> set[int]:
    holidays = panel.get("holidayDays")
    if not isinstance(holidays, list):
        return set()
    normalized = set()
    for day in holidays:
        try:
            normalized.add(int(day))
        except (TypeError, ValueError):
            continue
    return normalized


def block_panels(block: dict[str, object], result: dict[str, object]) -> list[dict[str, object]]:
    panels = block.get("monthPanels")
    if isinstance(panels, list) and panels:
        return [panel for panel in panels if isinstance(panel, dict)]
    return [
        {
            "label": "基準",
            "yearMonth": result.get("yearMonth", ""),
            "activeDays": result.get("activeDays"),
            "holidayDays": result.get("holidayDays", []),
            "days": result.get("days", []),
            "rows": block.get("rows", []),
        }
    ]


def block_title(block: dict[str, object], result: dict[str, object]) -> str:
    meta = block.get("meta") if isinstance(block.get("meta"), dict) else {}
    if block.get("kind") == "supplier":
        vend_cd = meta.get("vendCd", "")
        vend_name = meta.get("vendName") or vend_cd
        return " ".join(
            [
                f"仕入先：{vend_name}（{vend_cd}）",
                "入荷",
                f"品目番号：{meta.get('itemCd', '')}",
                f"手番：{meta.get('teban', '0')}",
                f"安全在庫：{meta.get('anzen', '0')}",
                f"手配区分：{meta.get('requestType', '')}",
                f"保管区：{meta.get('whCd', '')}",
                f"階層({meta.get('kaiso', 1)})",
            ]
        )

    cust_code = meta.get("custCode") or result.get("custCode", "")
    customer_name = meta.get("customerName") or result.get("customerName") or cust_code
    return " ".join(
        [
            f"得意先：{customer_name}（{cust_code}）",
            "出荷",
            f"得意先品目：{meta.get('custItem') or result.get('custItem', '')}",
            f"手番：{meta.get('teban', '0')}",
            f"安全在庫：{meta.get('anzen', '0')}",
        ]
    )


def build_gonenkukumi_workbook(result: dict[str, object]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "検索結果"

    condition_rows = [
        ("得意先", f"{result.get('customerName', '')}（{result.get('custCode', '')}）"),
        ("得意先品目", result.get("custItem", "")),
        ("品目任意変換値", result.get("optionChange", "")),
        ("検索年月", display_year_month(result.get("yearMonth", ""))),
        ("対象日付", display_date(result.get("asOfDate", ""))),
        ("内作品番", result.get("internalItemCd", "")),
    ]
    for index, (label, value) in enumerate(condition_rows):
        row = index + 1
        sheet.cell(row=row, column=1, value=label)
        sheet.cell(row=row, column=2, value=value)
        sheet.cell(row=row, column=1).fill = PatternFill("solid", fgColor="D9D9D9")
        sheet.cell(row=row, column=1).font = Font(name="MS PGothic", bold=True)
        sheet.cell(row=row, column=1).border = GRID_BORDER
        sheet.cell(row=row, column=2).border = GRID_BORDER

    current_row = 8
    day_headers = [f"{day:02d}" for day in range(1, 32)]
    outline_ranges: list[tuple[int, int, int, int]] = []
    month_header_border_rows: list[int] = []

    for block in result.get("blocks", []):
        colors = block_theme_colors(block)
        title = block_title(block, result)
        apply_range_fill(sheet, current_row, 1, 34, colors["bg"])
        sheet.cell(row=current_row, column=1, value=title)
        sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=34)
        outline_ranges.append((current_row, 1, current_row, 34))
        sheet.cell(row=current_row, column=1).font = Font(name="MS PGothic", bold=True)
        sheet.cell(row=current_row, column=1).alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[current_row].height = 24
        current_row += 1

        for panel in block_panels(block, result):
            active_days = panel_active_days(panel)
            holiday_days = panel_holiday_days(panel)
            panel_year_month = display_year_month(panel.get("yearMonth"))
            panel_label = str(panel.get("label") or "")
            panel_title = panel_year_month if panel_label in {"", "基準"} else f"{panel_year_month}（{panel_label}）"
            sheet.cell(row=current_row, column=1, value=panel_title)
            sheet.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=34)
            apply_range_fill(sheet, current_row, 1, 34, colors["bg"])
            panel_start_row = current_row
            outline_ranges.append((current_row, 1, current_row, 34))
            month_header_border_rows.append(current_row)
            sheet.cell(row=current_row, column=1).font = Font(name="MS PGothic", bold=True)
            current_row += 1

            headers = ["区分", "行計", "前月残", *day_headers]
            for col, value in enumerate(headers, start=1):
                day_no = col - 3
                is_pad_day = col >= 4 and day_no > active_days
                cell = sheet.cell(row=current_row, column=col, value="" if is_pad_day else value)
                cell.fill = PatternFill("solid", fgColor=PAD_DAY_COLOR if is_pad_day else colors["bg"])
                cell.font = Font(name="MS PGothic", bold=True)
                cell.border = GRID_BORDER
                cell.alignment = Alignment(horizontal="center")
            current_row += 1

            for row in panel.get("rows", []):
                values = row.get("values", [])
                row_total = row.get("total", sum(value for value in values if isinstance(value, Number)))
                sheet.cell(row=current_row, column=1, value=row.get("item", ""))
                sheet.cell(row=current_row, column=2, value=display_number(row_total))
                sheet.cell(row=current_row, column=3, value=display_number(row.get("balance")))
                apply_range_fill(sheet, current_row, 1, 34, colors["veryLightBg"])
                for col, value in enumerate(values[:31], start=4):
                    day_no = col - 3
                    cell = sheet.cell(row=current_row, column=col, value=display_number(value))
                    if day_no > active_days:
                        cell.value = None
                        cell.fill = PatternFill("solid", fgColor=PAD_DAY_COLOR)
                    elif day_no in holiday_days:
                        cell.fill = PatternFill("solid", fgColor=colors["bg"])
                current_row += 1
            outline_ranges.append((panel_start_row, 1, current_row - 1, 34))
        current_row += 1

    for row in sheet.iter_rows():
        for cell in row:
            cell.font = Font(name="MS PGothic", size=10)
            if cell.value is not None or cell.fill.fill_type is not None:
                cell.border = GRID_BORDER
            if isinstance(cell.value, Number):
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right")

    for outline_range in outline_ranges:
        apply_outer_border(sheet, *outline_range)
    for row in month_header_border_rows:
        apply_horizontal_border(sheet, row, 1, 34)

    fit_column_widths(sheet)
    sheet.freeze_panes = "A8"

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
