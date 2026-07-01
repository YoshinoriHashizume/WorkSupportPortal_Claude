from io import BytesIO

from openpyxl import load_workbook

from applications.gonenkukumi.domain.excel_export import build_gonenkukumi_workbook


def test_excel_export_hides_zero_values_and_keeps_number_format():
    workbook_bytes = build_gonenkukumi_workbook(
        {
            "customerName": "得意先A",
            "custCode": "101",
            "custItem": "ITEM-001",
            "optionChange": "*",
            "yearMonth": "2026-05",
            "asOfDate": "2026-05-21",
            "activeDays": 30,
            "holidayDays": [2],
            "internalItemCd": "NAISAK-001",
            "blocks": [
                {
                    "label": "得意先 101",
                    "theme": {
                        "bg": "#dcfce7",
                        "lightBg": "#f0fdf4",
                        "veryLightBg": "#f7fef9",
                        "outerBorder": "#22c55e",
                    },
                    "meta": {
                        "custCode": "101",
                        "customerName": "得意先A",
                        "custItem": "ITEM-001",
                        "teban": 0,
                        "anzen": 0,
                    },
                    "rows": [
                        {"item": "内示受注", "values": [0, 1234, *([0] * 29)], "total": 1234},
                        {"item": "確定受注", "values": [0 for _ in range(31)], "total": 0},
                    ],
                }
            ],
        }
    )

    workbook = load_workbook(BytesIO(workbook_bytes))
    sheet = workbook.active
    values = [cell.value for row in sheet.iter_rows() for cell in row]

    assert sheet.title == "検索結果"
    assert sheet["A1"].value == "得意先"
    assert sheet["B1"].value == "得意先A（101）"
    assert sheet["A4"].value == "検索年月"
    assert sheet["B4"].value == "2026/05"
    assert sheet["A5"].value == "対象日付"
    assert sheet["B5"].value == "2026/05/21"
    assert sheet["A6"].value == "内作品番"
    assert sheet["B6"].value == "NAISAK-001"
    assert sheet.column_dimensions["A"].bestFit is True
    assert sheet.column_dimensions["B"].bestFit is True
    assert sheet.column_dimensions["A"].width >= 10
    assert sheet.column_dimensions["B"].width >= 14
    assert sheet.freeze_panes == "A8"
    assert sheet["A8"].value == "得意先：得意先A（101） 出荷 得意先品目：ITEM-001 手番：0 安全在庫：0"
    assert "A8:AH8" in {str(range_ref) for range_ref in sheet.merged_cells.ranges}
    assert sheet["A8"].border.left.style == "medium"
    assert sheet["A8"].border.top.style == "medium"
    assert sheet["A8"].border.right.style == "medium"
    assert sheet["A8"].border.bottom.style == "medium"
    assert sheet["A9"].border.top.style == "medium"
    assert sheet["A9"].border.bottom.style == "thin"
    assert sheet["A9"].border.bottom.color.rgb.endswith("334155")
    assert sheet["A10"].border.top.style == "thin"
    assert sheet["A10"].border.top.color.rgb.endswith("334155")
    assert sheet["B10"].border.left.style == "thin"
    assert sheet["A12"].border.bottom.style == "medium"
    assert sheet["A13"].border.top.style is None
    assert sheet["A13"].border.bottom.style is None
    assert sheet["A8"].fill.fgColor.rgb.endswith("DCFCE7")
    assert sheet["A10"].fill.fgColor.rgb.endswith("DCFCE7")
    assert sheet["A11"].fill.fgColor.rgb.endswith("F7FEF9")
    assert sheet["E11"].fill.fgColor.rgb.endswith("DCFCE7")
    assert sheet["AH10"].value is None
    assert sheet["AH10"].fill.fgColor.rgb.endswith("B8C4D0")
    assert sheet["AH11"].value is None
    assert sheet["AH11"].fill.fgColor.rgb.endswith("B8C4D0")
    assert "対象" not in values
    assert sheet["A10"].value == "区分"
    assert sheet["B10"].value == "行計"
    assert sheet["C10"].value == "前月残"
    assert 0 not in values
    assert 1234 in values

    formatted_cell = next(cell for row in sheet.iter_rows() for cell in row if cell.value == 1234)
    assert formatted_cell.number_format == "#,##0"


def test_excel_export_outputs_month_panels_and_previous_balance_column():
    workbook_bytes = build_gonenkukumi_workbook(
        {
            "customerName": "得意先A",
            "custCode": "101",
            "custItem": "ITEM-001",
            "optionChange": "*",
            "yearMonth": "2026-05",
            "internalItemCd": "NAISAK-001",
            "months": [
                {"yearMonth": "2026-04", "label": "前月"},
                {"yearMonth": "2026-05", "label": "基準"},
            ],
            "blocks": [
                {
                    "label": "得意先 101",
                    "monthPanels": [
                        {
                            "yearMonth": "2026-04",
                            "label": "前月",
                            "rows": [{"item": "確定受注", "values": [100], "total": 150, "balance": 50}],
                        },
                        {
                            "yearMonth": "2026-05",
                            "label": "基準",
                            "rows": [{"item": "確定受注", "values": [200], "total": 260, "balance": 60}],
                        },
                    ],
                },
                {
                    "kind": "supplier",
                    "label": "仕入先 9200",
                    "theme": {
                        "bg": "hsl(200.000 55% 92%)",
                        "lightBg": "hsl(200.000 45% 96%)",
                        "veryLightBg": "hsl(200.000 25% 98%)",
                    },
                    "meta": {
                        "vendCd": "9200",
                        "vendName": "テクノス工場",
                        "itemCd": "235677-0050-T-9200",
                        "teban": 3,
                        "anzen": 0,
                        "requestType": "かんばん",
                        "whCd": "M80",
                        "kaiso": 1,
                    },
                    "monthPanels": [
                        {
                            "yearMonth": "2026-05",
                            "label": "基準",
                            "rows": [{"item": "月初発注", "values": [300], "total": 300, "balance": 0}],
                        },
                    ],
                }
            ],
        }
    )

    workbook = load_workbook(BytesIO(workbook_bytes))
    sheet = workbook.active
    values = [cell.value for row in sheet.iter_rows() for cell in row]

    assert "2026/04（前月）" in values
    assert "2026/05" in values
    assert "2026/05（基準）" not in values
    assert sheet["B4"].value == "2026/05"
    assert "2026-04, 2026-05" not in values
    assert "前月残" in values
    assert (
        "仕入先：テクノス工場（9200） 入荷 品目番号：235677-0050-T-9200 "
        "手番：3 安全在庫：0 手配区分：かんばん 保管区：M80 階層(1)"
    ) in values
    assert 50 in values
    assert 60 in values
